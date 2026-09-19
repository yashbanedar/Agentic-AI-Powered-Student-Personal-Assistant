import os
import pypdf
import docx
import pptx
from typing import Dict, Any, List
from app.core.rag import chunk_text

class DocumentAgent:
    """
    Document Agent: Extracts text and structure from PDF, PPTX, DOCX, TXT.
    Generates chunks with page/slide metadata for RAG retrieval.
    """
    
    @classmethod
    def extract_text(cls, file_path: str, file_type: str) -> Dict[str, Any]:
        file_type = file_type.lower().strip(".")
        extracted_text = ""
        sections = []
        
        try:
            if file_type == "pdf":
                reader = pypdf.PdfReader(file_path)
                for page_idx, page in enumerate(reader.pages):
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        extracted_text += f"\n--- Page {page_idx + 1} ---\n" + page_text
                        sections.append({
                            "page": page_idx + 1,
                            "type": "page",
                            "content": page_text.strip()
                        })
                        
            elif file_type in ["docx", "doc"]:
                doc = docx.Document(file_path)
                for para_idx, para in enumerate(doc.paragraphs):
                    text = para.text.strip()
                    if text:
                        extracted_text += "\n" + text
                        sections.append({
                            "paragraph": para_idx + 1,
                            "type": "paragraph",
                            "content": text
                        })
                        
            elif file_type in ["pptx", "ppt"]:
                prs = pptx.Presentation(file_path)
                for slide_idx, slide in enumerate(prs.slides):
                    slide_texts = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_texts.append(shape.text.strip())
                    combined_slide = " \n".join(slide_texts)
                    if combined_slide:
                        extracted_text += f"\n--- Slide {slide_idx + 1} ---\n" + combined_slide
                        sections.append({
                            "slide": slide_idx + 1,
                            "type": "slide",
                            "content": combined_slide
                        })
            else: # txt, md, note
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    extracted_text = f.read()
                sections.append({
                    "type": "full_text",
                    "content": extracted_text
                })
        except Exception as e:
            extracted_text = f"Error extracting text from {file_type} file: {str(e)}"
            
        chunks = chunk_text(extracted_text, chunk_size=700, chunk_overlap=120)
        
        chunk_objects = []
        for idx, ch in enumerate(chunks):
            # associate with page/slide if possible
            page_hint = None
            if "--- Page " in ch:
                parts = ch.split("--- Page ")
                if len(parts) > 1:
                    page_hint = parts[1].split(" ---")[0]
            elif "--- Slide " in ch:
                parts = ch.split("--- Slide ")
                if len(parts) > 1:
                    page_hint = parts[1].split(" ---")[0]
                    
            chunk_objects.append({
                "chunk_index": idx,
                "content": ch,
                "metadata_json": {
                    "page_or_slide": page_hint or f"Chunk {idx + 1}",
                    "word_count": len(ch.split())
                }
            })
            
        return {
            "extracted_text": extracted_text,
            "chunk_count": len(chunk_objects),
            "chunks": chunk_objects,
            "summary": f"Extracted {len(extracted_text.split())} words across {len(chunk_objects)} semantic chunks."
        }
